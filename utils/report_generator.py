"""
utils/report_generator.py - Generador de informes ejecutivos en PDF (ReportLab) y Excel (openpyxl).
Diseñado bajo la estética "clinical blueprint on frosted paper" (monocromático con acento #e7000b).
"""

import io
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)

def generate_excel_report(df_factores: pd.DataFrame, df_recomendaciones: pd.DataFrame) -> bytes:
    """Genera un archivo Excel profesional con múltiples hojas y formato clínico monocromático."""
    output = io.BytesIO()
    wb = Workbook()

    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=13, bold=True, color="0A0A0A")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="737373")
    regular_font = Font(name="Segoe UI", size=9.5)
    bold_cell_font = Font(name="Segoe UI", size=9.5, bold=True)

    header_fill_dark = PatternFill(start_color="0A0A0A", end_color="0A0A0A", fill_type="solid")
    header_fill_soft = PatternFill(start_color="171717", end_color="171717", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E5E5E5'),
        right=Side(style='thin', color='E5E5E5'),
        top=Side(style='thin', color='E5E5E5'),
        bottom=Side(style='thin', color='E5E5E5')
    )

    # ------------------ HOJA 1: RESUMEN EJECUTIVO ------------------
    ws1 = wb.active
    ws1.title = "Resumen Ejecutivo"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:G1")
    ws1["A1"] = "SMURFIT WESTROCK COLOMBIA — LECTURA DEL ENTORNO PESTEL"
    ws1["A1"].font = title_font
    ws1["A1"].alignment = Alignment(vertical="center")

    ws1.merge_cells("A2:G2")
    ws1["A2"] = f"Complejo Industrial Yumbo, Valle del Cauca · Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws1["A2"].font = subtitle_font

    kpis = [
        ("Total de Factores Monitoreados", len(df_factores)),
        ("Riesgos Identificados", len(df_factores[df_factores["tipo"] == "Riesgo"])),
        ("Oportunidades Identificadas", len(df_factores[df_factores["tipo"] == "Oportunidad"])),
        ("Riesgos Críticos (Severidad >= 16)", len(df_factores[(df_factores["tipo"] == "Riesgo") & (df_factores["puntuacion_severidad"] >= 16)])),
        ("Oportunidades Clave (Severidad >= 16)", len(df_factores[(df_factores["tipo"] == "Oportunidad") & (df_factores["puntuacion_severidad"] >= 16)])),
        ("Total Decisiones y Recomendaciones", len(df_recomendaciones))
    ]

    ws1.append([])
    ws1.append(["Indicador Estratégico", "Valor"])
    ws1["A4"].font = header_font
    ws1["A4"].fill = header_fill_dark
    ws1["B4"].font = header_font
    ws1["B4"].fill = header_fill_dark

    for row_idx, (kpi, val) in enumerate(kpis, start=5):
        ws1.append([kpi, val])
        ws1[f"A{row_idx}"].font = regular_font
        ws1[f"B{row_idx}"].font = bold_cell_font
        ws1[f"A{row_idx}"].border = thin_border
        ws1[f"B{row_idx}"].border = thin_border

    ws1.append([])
    pestel_start = ws1.max_row + 1
    ws1.append(["Dimensión PESTEL", "Total Factores", "Riesgos", "Oportunidades", "Impacto Promedio"])
    for col_letter in ["A", "B", "C", "D", "E"]:
        ws1[f"{col_letter}{pestel_start}"].font = header_font
        ws1[f"{col_letter}{pestel_start}"].fill = header_fill_soft

    for cat in ["Político", "Económico", "Social", "Tecnológico", "Ecológico", "Legal"]:
        sub = df_factores[df_factores["categoria_pestel"] == cat]
        t = len(sub)
        r = len(sub[sub["tipo"] == "Riesgo"])
        o = len(sub[sub["tipo"] == "Oportunidad"])
        imp_prom = round(sub["impacto"].mean(), 1) if t > 0 else 0
        ws1.append([cat, t, r, o, imp_prom])
        curr_row = ws1.max_row
        for col_letter in ["A", "B", "C", "D", "E"]:
            ws1[f"{col_letter}{curr_row}"].font = regular_font
            ws1[f"{col_letter}{curr_row}"].border = thin_border

    # ------------------ HOJA 2: FACTORES PESTEL ------------------
    ws2 = wb.create_sheet(title="Factores PESTEL")
    ws2.views.sheetView[0].showGridLines = True
    cols_factores = [
        "Código", "Categoría PESTEL", "Tipo", "Ámbito", "Título del Factor",
        "Probabilidad (1-5)", "Impacto (1-5)", "Severidad (1-25)", "Urgencia",
        "Efecto en Smurfit WestRock Yumbo", "Fuente", "URL Fuente", "Fecha Publicación",
        "Fecha Consulta", "Estado"
    ]
    ws2.append(cols_factores)
    for col_num in range(1, len(cols_factores) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill_dark
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for _, f in df_factores.iterrows():
        ws2.append([
            f["codigo"], f["categoria_pestel"], f["tipo"], f.get("ambito", "Nacional"),
            f["titulo"], f["probabilidad"], f["impacto"], f["puntuacion_severidad"],
            f["urgencia"], f["efecto_empresa"], f["fuente_nombre"], f["fuente_url"],
            f["fecha_publicacion"], f["fecha_consulta"], f["estado"]
        ])
        curr_row = ws2.max_row
        for c_idx in range(1, len(cols_factores) + 1):
            ws2.cell(row=curr_row, column=c_idx).font = regular_font
            ws2.cell(row=curr_row, column=c_idx).border = thin_border

    # ------------------ HOJA 3: RECOMENDACIONES ------------------
    ws3 = wb.create_sheet(title="Plan de Acción")
    ws3.views.sheetView[0].showGridLines = True
    cols_rec = [
        "Cód. Recomendación", "Factor Vinculado", "Título de la Recomendación",
        "Acción Estratégica Sugerida", "Evidencia de Respaldo", "Prioridad",
        "Área Responsable", "Plazo Estimado", "Estado"
    ]
    ws3.append(cols_rec)
    for col_num in range(1, len(cols_rec) + 1):
        cell = ws3.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill_soft
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if not df_recomendaciones.empty:
        for _, r in df_recomendaciones.iterrows():
            factor_info = f"[{r.get('factor_codigo', '')}] {r.get('factor_titulo', '')}"
            ws3.append([
                r["codigo_recomendacion"], factor_info, r["titulo"],
                r["accion_sugerida"], r["evidencia"], r["prioridad"],
                r["area_responsable"], r["plazo"], r["estado"]
            ])
            curr_row = ws3.max_row
            for c_idx in range(1, len(cols_rec) + 1):
                ws3.cell(row=curr_row, column=c_idx).font = regular_font
                ws3.cell(row=curr_row, column=c_idx).border = thin_border

    # Autoajuste de anchos de columna
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if '\n' in val_str:
                    val_str = val_str.split('\n')[0]
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

    wb.save(output)
    output.seek(0)
    return output.getvalue()

def generate_pdf_report(df_factores: pd.DataFrame, df_recomendaciones: pd.DataFrame) -> bytes:
    """Genera un informe ejecutivo formal en PDF con tipografía geométrica neutra."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0a0a0a')
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#737373')
    )
    section_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#0a0a0a'),
        spaceBefore=10,
        spaceAfter=5
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#171717')
    )
    table_cell = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#0a0a0a')
    )
    table_cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=table_cell,
        fontName='Helvetica-Bold'
    )

    story = []

    # 1. ENCABEZADO MONOCROMÁTICO
    story.append(Paragraph("SMURFIT WESTROCK COLOMBIA", title_style))
    story.append(Paragraph("Informe de Lectura del Entorno Estratégico (PESTEL)", ParagraphStyle(
        'SubT', parent=styles['Heading3'], fontName='Helvetica-Bold', fontSize=11, textColor=colors.HexColor('#171717')
    )))
    fecha_hoy = datetime.now().strftime("%d de %B de %Y")
    story.append(Paragraph(f"Complejo Industrial Yumbo (Valle del Cauca) · Emisión: {fecha_hoy} · Uso Interno", subtitle_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e5e5e5'), spaceBefore=2, spaceAfter=8))

    # 2. RESUMEN EJECUTIVO
    story.append(Paragraph("1. Síntesis Ejecutiva", section_style))
    total_f = len(df_factores)
    riesgos = df_factores[df_factores["tipo"] == "Riesgo"]
    oportunidades = df_factores[df_factores["tipo"] == "Oportunidad"]
    criticos = riesgos[riesgos["puntuacion_severidad"] >= 16]
    estrategicas = oportunidades[oportunidades["puntuacion_severidad"] >= 16]

    resumen_text = (
        f"Consolidado de vigilancia del entorno para <b>Smurfit WestRock Colombia</b> focalizado en las operaciones "
        f"del complejo papelero y de empaques en Yumbo y sus núcleos forestales. Se han evaluado un total de "
        f"<b>{total_f} factores estratégicos</b> ({len(riesgos)} riesgos y {len(oportunidades)} oportunidades). "
        f"Se destacan <b>{len(criticos)} riesgos críticos</b> que requieren planes de contingencia (principalmente en "
        f"el corredor vial Panamericano y la cuenca del Río Cauca) y <b>{len(estrategicas)} oportunidades de alto impacto</b> "
        f"asociadas a la Ley 2232 de plásticos de un solo uso y el crecimiento agroexportador."
    )
    story.append(Paragraph(resumen_text, body_style))
    story.append(Spacer(1, 8))

    # Métricas en caja sutil
    kpi_table_data = [
        [
            Paragraph("<b>Total Factores:</b> " + str(total_f), table_cell_bold),
            Paragraph("<b>Riesgos Críticos:</b> <font color='#e7000b'>" + str(len(criticos)) + "</font>", table_cell_bold),
            Paragraph("<b>Oportunidades Clave:</b> " + str(len(estrategicas)), table_cell_bold),
            Paragraph("<b>Recomendaciones:</b> " + str(len(df_recomendaciones)), table_cell_bold)
        ]
    ]
    kpi_table = Table(kpi_table_data, colWidths=[135, 135, 135, 135])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#fafafa')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e5e5e5')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e5e5')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 10))

    # 3. DISTRIBUCIÓN PESTEL
    story.append(Paragraph("2. Balance por Dimensión PESTEL", section_style))
    pestel_table_data = [
        [
            Paragraph("<b>Dimensión</b>", table_cell_bold),
            Paragraph("<b>Total</b>", table_cell_bold),
            Paragraph("<b>Riesgos</b>", table_cell_bold),
            Paragraph("<b>Oportunidades</b>", table_cell_bold),
            Paragraph("<b>Impacto Promedio</b>", table_cell_bold),
            Paragraph("<b>Factor Principal</b>", table_cell_bold)
        ]
    ]

    for cat in ["Político", "Económico", "Social", "Tecnológico", "Ecológico", "Legal"]:
        sub = df_factores[df_factores["categoria_pestel"] == cat]
        t = len(sub)
        r = len(sub[sub["tipo"] == "Riesgo"])
        o = len(sub[sub["tipo"] == "Oportunidad"])
        imp_prom = f"{round(sub['impacto'].mean(), 1)} / 5" if t > 0 else "-"
        top_factor = sub.sort_values(by="puntuacion_severidad", ascending=False).iloc[0]["titulo"][:50] + "..." if t > 0 else "-"

        pestel_table_data.append([
            Paragraph(cat, table_cell_bold),
            Paragraph(str(t), table_cell),
            Paragraph(str(r), table_cell),
            Paragraph(str(o), table_cell),
            Paragraph(imp_prom, table_cell),
            Paragraph(top_factor, table_cell)
        ])

    pestel_table = Table(pestel_table_data, colWidths=[75, 40, 45, 55, 75, 250])
    pestel_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a0a0a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e5e5')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#fafafa')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(pestel_table)
    story.append(Spacer(1, 10))

    # 4. FACTORES CRÍTICOS Y ESTRATÉGICOS
    story.append(Paragraph("3. Factores de Mayor Severidad (P x I >= 16)", section_style))
    top_factors = df_factores[df_factores["puntuacion_severidad"] >= 16].sort_values(
        by=["tipo", "puntuacion_severidad"], ascending=[False, False]
    )

    factor_table_data = [
        [
            Paragraph("<b>Cód</b>", table_cell_bold),
            Paragraph("<b>Tipo</b>", table_cell_bold),
            Paragraph("<b>Dimensión</b>", table_cell_bold),
            Paragraph("<b>Título & Efecto en Yumbo</b>", table_cell_bold),
            Paragraph("<b>P x I</b>", table_cell_bold),
            Paragraph("<b>Urgencia</b>", table_cell_bold),
            Paragraph("<b>Fuente</b>", table_cell_bold)
        ]
    ]

    for _, row in top_factors.iterrows():
        is_crit_risk = row["tipo"] == "Riesgo" and row["puntuacion_severidad"] >= 17
        tipo_color = "#e7000b" if is_crit_risk else "#171717"
        tipo_cell = f"<font color='{tipo_color}'><b>{row['tipo']}</b></font>"
        desc_combo = f"<b>{row['titulo']}</b><br/>{row['efecto_empresa'][:100]}..."
        pxi = f"{row['probabilidad']}x{row['impacto']}=<b>{row['puntuacion_severidad']}</b>"
        fuente_info = f"{row['fuente_nombre'][:22]} ({row['fecha_publicacion']})"

        factor_table_data.append([
            Paragraph(row["codigo"], table_cell_bold),
            Paragraph(tipo_cell, table_cell),
            Paragraph(row["categoria_pestel"], table_cell),
            Paragraph(desc_combo, table_cell),
            Paragraph(pxi, table_cell),
            Paragraph(row["urgencia"][:14], table_cell),
            Paragraph(fuente_info, table_cell)
        ])

    factor_table = Table(factor_table_data, colWidths=[45, 55, 60, 210, 45, 65, 60])
    factor_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#171717')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e5e5')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#fafafa')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(factor_table)
    story.append(Spacer(1, 10))

    # 5. RECOMENDACIONES
    story.append(Paragraph("4. Plan de Acción y Decisiones Estratégicas", section_style))
    if not df_recomendaciones.empty:
        rec_table_data = [
            [
                Paragraph("<b>Código</b>", table_cell_bold),
                Paragraph("<b>Recomendación & Acción Sugerida</b>", table_cell_bold),
                Paragraph("<b>Prioridad</b>", table_cell_bold),
                Paragraph("<b>Área Responsable</b>", table_cell_bold),
                Paragraph("<b>Plazo</b>", table_cell_bold),
                Paragraph("<b>Estado</b>", table_cell_bold)
            ]
        ]

        for _, rec in df_recomendaciones.iterrows():
            content = f"<b>{rec['titulo']}</b><br/>{rec['accion_sugerida'][:110]}..."
            prio_color = "#e7000b" if rec['prioridad'] == "Alta" else "#171717"
            prio_cell = f"<font color='{prio_color}'><b>{rec['prioridad']}</b></font>"

            rec_table_data.append([
                Paragraph(rec["codigo_recomendacion"], table_cell_bold),
                Paragraph(content, table_cell),
                Paragraph(prio_cell, table_cell),
                Paragraph(rec["area_responsable"], table_cell),
                Paragraph(rec["plazo"][:15], table_cell),
                Paragraph(rec["estado"], table_cell)
            ])

        rec_table = Table(rec_table_data, colWidths=[55, 230, 45, 105, 55, 50])
        rec_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a0a0a')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e5e5')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#fafafa')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(rec_table)

    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#e5e5e5'), spaceBefore=4, spaceAfter=6))
    story.append(Paragraph(
        "Documento estratégico confidencial · Sistema de Inteligencia PESTEL · Smurfit WestRock Colombia (Yumbo)",
        ParagraphStyle('Footer', parent=styles['Normal'], fontSize=7.5, textColor=colors.HexColor('#737373'), alignment=1)
    ))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
